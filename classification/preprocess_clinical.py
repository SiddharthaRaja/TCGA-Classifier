#!/usr/bin/env python3
"""
TCGA-BRCA Clinical Data Preprocessor (Fixed)
==============================================
Reads per-patient Clinical Supplement XMLs downloaded by tcga_brca_downloader.py,
extracts key fields for ML/DL, and outputs:
  - clinical_preprocessed.csv   (one row per patient, ML-ready)
  - clinical_preprocessed.xlsx  (same data, formatted with a summary sheet)

Requirements:
  pip install pandas openpyxl lxml
"""

import re
import warnings
from pathlib import Path

import pandas as pd
from lxml import etree
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

# ── Configuration ─────────────────────────────────────────────────────────────
DATA_DIR          = Path("TCGA_BRCA_data")
OUTPUT_CSV        = Path("clinical_preprocessed.csv")
OUTPUT_XLSX       = Path("clinical_preprocessed.xlsx")
MISSING_THRESHOLD = 0.60   # drop column if > 60% values are missing
# ─────────────────────────────────────────────────────────────────────────────


def find_clinical_xmls(data_dir: Path):
    """
    Returns only the main clinical XML per patient:
      nationwidechildrens.org_clinical.TCGA-*.xml
    Ignores biospecimen, ssf, and other supplementary files.
    """
    found = []
    for xml_path in sorted(data_dir.glob("*/Clinical/*.xml")):
        name = xml_path.name.lower()
        if "_clinical." in name and "biospecimen" not in name and "ssf" not in name:
            found.append(xml_path)
    return found


def _find_tag(root, *local_names):
    """Search entire XML tree for the first matching local tag name."""
    for el in root.iter():
        local = el.tag.split("}")[-1] if "}" in el.tag else el.tag
        if local in local_names:
            if el.text and el.text.strip():
                return el.text.strip()
    return None


def parse_clinical_xml(xml_path: Path) -> dict:
    tree = etree.parse(str(xml_path))
    root = tree.getroot()

    rec = {"patient_id": xml_path.parent.parent.name}

    # Demographics
    rec["age_at_diagnosis"] = _find_tag(root,
        "age_at_initial_pathologic_diagnosis", "age_at_diagnosis")
    rec["gender"]    = _find_tag(root, "gender")
    rec["race"]      = _find_tag(root, "race")
    rec["ethnicity"] = _find_tag(root, "ethnicity")

    # Tumor info
    rec["tumor_stage"] = _find_tag(root,
        "clinical_stage", "pathologic_stage", "tumor_stage", "stage_event")
    rec["tumor_grade"] = _find_tag(root,
        "neoplasm_histologic_grade", "histologic_grade", "tumor_grade")
    rec["histological_type"] = _find_tag(root, "histological_type", "histology")
    rec["tumor_site"]        = _find_tag(root,
        "tumor_tissue_site", "primary_site", "tissue_or_organ_of_origin")

    # Receptor status
    rec["er_status"]   = _find_tag(root,
        "breast_carcinoma_estrogen_receptor_status", "er_status_by_ihc", "er_status")
    rec["pr_status"]   = _find_tag(root,
        "breast_carcinoma_progesterone_receptor_status", "pr_status_by_ihc", "pr_status")
    rec["her2_status"] = _find_tag(root,
        "lab_proc_her2_neu_immunohistochemistry_receptor_status",
        "her2_status_by_ihc", "her2_fish_status", "her2_status")

    # Survival / follow-up
    rec["vital_status"]             = _find_tag(root, "vital_status")
    rec["days_to_death"]            = _find_tag(root, "days_to_death")
    rec["days_to_last_followup"]    = _find_tag(root,
        "days_to_last_followup", "days_to_last_follow_up")
    rec["days_to_last_known_alive"] = _find_tag(root, "days_to_last_known_alive")

    # Treatment
    rec["radiation_therapy"]      = _find_tag(root,
        "radiation_therapy", "radiation_treatment_ongoing")
    rec["pharmaceutical_therapy"] = _find_tag(root,
        "targeted_molecular_therapy", "pharmaceutical_therapy")
    rec["treatment_outcome"]      = _find_tag(root,
        "treatment_outcome_first_course", "primary_therapy_outcome_success")

    return rec


_NOT_AVAILABLE = {
    "[not available]", "[not applicable]", "[unknown]", "[discrepancy]",
    "not reported", "na", "n/a", "unknown", "null", "none", ""
}

def normalise(val):
    if val is None:
        return None
    v = str(val).strip().lower()
    return None if v in _NOT_AVAILABLE else v


_STAGE_MAP = {
    "i": 1, "stage i": 1, "stage ia": 1, "stage ib": 1,
    "ii": 2, "stage ii": 2, "stage iia": 2, "stage iib": 2,
    "iii": 3, "stage iii": 3, "stage iiia": 3, "stage iiib": 3, "stage iiic": 3,
    "iv": 4, "stage iv": 4,
}
_YES = {"yes", "positive", "administered", "true", "1"}
_NO  = {"no", "negative", "not administered", "false", "0"}

def encode_stage(val):
    if val is None: return None
    return _STAGE_MAP.get(str(val).lower().strip())

def encode_grade(val):
    if val is None: return None
    m = re.search(r"\b([1-4])\b", str(val))
    return int(m.group(1)) if m else None

def encode_binary(val):
    if val is None: return None
    v = str(val).strip().lower()
    if v in _YES: return 1
    if v in _NO:  return 0
    return None


def build_dataframe(records: list) -> pd.DataFrame:
    df = pd.DataFrame(records)

    for col in df.columns:
        if col != "patient_id":
            df[col] = df[col].apply(normalise)

    for col in ["age_at_diagnosis", "days_to_death",
                "days_to_last_followup", "days_to_last_known_alive"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    df["tumor_stage_numeric"] = df["tumor_stage"].apply(encode_stage)
    df["tumor_grade_numeric"] = df["tumor_grade"].apply(encode_grade)
    df["vital_status_binary"] = df["vital_status"].apply(encode_binary)
    df["er_status_binary"]    = df["er_status"].apply(encode_binary)
    df["pr_status_binary"]    = df["pr_status"].apply(encode_binary)
    df["her2_status_binary"]  = df["her2_status"].apply(encode_binary)
    df["radiation_binary"]    = df["radiation_therapy"].apply(encode_binary)
    df["pharma_binary"]       = df["pharmaceutical_therapy"].apply(encode_binary)

    for col in ["gender", "race", "ethnicity", "histological_type",
                "tumor_site", "treatment_outcome"]:
        if col in df.columns and df[col].notna().any():
            dummies = pd.get_dummies(df[col], prefix=col, dummy_na=False)
            df = pd.concat([df, dummies], axis=1)

    return df


def drop_high_missing(df: pd.DataFrame, threshold=MISSING_THRESHOLD) -> pd.DataFrame:
    missing_rate = df.isna().mean()
    to_drop = missing_rate[missing_rate > threshold].index.tolist()
    if to_drop:
        print(f"\n  Dropping {len(to_drop)} column(s) with >{threshold*100:.0f}% missing:")
        for c in to_drop:
            print(f"    - {c}  ({missing_rate[c]*100:.0f}% missing)")
    return df.drop(columns=to_drop)


def impute_numeric(df: pd.DataFrame) -> pd.DataFrame:
    for col in df.select_dtypes(include="number").columns:
        if df[col].isna().any():
            df[col] = df[col].fillna(df[col].median())
    return df


def write_xlsx(df: pd.DataFrame, path: Path):
    summary_data = {
        "Metric": [
            "Total patients", "Total features", "Avg missing rate",
            "Numeric features", "Object/text features", "One-hot encoded columns",
        ],
        "Value": [
            len(df),
            len(df.columns),
            f"{df.isna().mean().mean()*100:.1f}%",
            len(df.select_dtypes(include="number").columns),
            len(df.select_dtypes(include="object").columns),
            len([c for c in df.columns if any(
                c.startswith(p + "_") for p in
                ["gender", "race", "ethnicity", "histological_type",
                 "tumor_site", "treatment_outcome"])]),
        ]
    }

    with pd.ExcelWriter(str(path), engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Clinical Data", index=False)
        pd.DataFrame(summary_data).to_excel(writer, sheet_name="Summary", index=False)

    wb = load_workbook(str(path))
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    header_fill = PatternFill("solid", start_color="1F4E79")
    alt_fill    = PatternFill("solid", start_color="D6E4F0")
    center      = Alignment(horizontal="center", vertical="center")
    thin        = Side(style="thin", color="CCCCCC")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            fill = alt_fill if row_idx % 2 == 0 else PatternFill()
            for cell in row:
                cell.fill = fill
                cell.alignment = center
                cell.border = border
        for col in ws.columns:
            max_len = max(
                (len(str(cell.value)) if cell.value is not None else 0)
                for cell in col
            )
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 45)
        ws.freeze_panes = "A2"

    wb.save(str(path))


def main():
    print("=" * 60)
    print("  TCGA-BRCA Clinical Preprocessor")
    print("=" * 60)

    xml_files = find_clinical_xmls(DATA_DIR)

    if not xml_files:
        print(f"\n  ERROR: No clinical XML files found under {DATA_DIR}/")
        print("    Run tcga_brca_downloader.py first.")
        return

    print(f"\n  Found {len(xml_files)} clinical XML file(s):")
    for xf in xml_files:
        print(f"    - {xf}")

    print("\n[1/5] Parsing XML files ...")
    records = []
    for xf in xml_files:
        try:
            rec = parse_clinical_xml(xf)
            records.append(rec)
            print(f"    OK: {rec['patient_id']}")
        except Exception as e:
            print(f"    FAIL: {xf.name} -> {e}")

    if not records:
        print("\n  No records parsed. Exiting.")
        return

    print("\n[2/5] Building DataFrame ...")
    df = build_dataframe(records)
    print(f"    Shape before cleaning: {df.shape}")

    print(f"\n[3/5] Dropping columns with >{MISSING_THRESHOLD*100:.0f}% missing ...")
    df = drop_high_missing(df, MISSING_THRESHOLD)

    print("\n[4/5] Imputing remaining numeric NaNs with column median ...")
    df = impute_numeric(df)
    print(f"    Shape after cleaning: {df.shape}")

    print("\n[5/5] Exporting outputs ...")
    df.to_csv(OUTPUT_CSV, index=False)
    print(f"    CSV  saved: {OUTPUT_CSV.resolve()}")
    write_xlsx(df, OUTPUT_XLSX)
    print(f"    XLSX saved: {OUTPUT_XLSX.resolve()}")

    print("\n" + "=" * 60)
    print("  Final columns:")
    print("=" * 60)
    for col in df.columns:
        missing = df[col].isna().sum()
        print(f"  {col:<45} {str(df[col].dtype):<12}  missing={missing}")

    print(f"\n  Done!  {len(df)} patient(s) x {len(df.columns)} features")
    print(f"  Output folder: {Path('.').resolve()}")


if __name__ == "__main__":
    main()
