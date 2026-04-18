#!/usr/bin/env python3
"""
TCGA-BRCA Clinical Preprocessing - Final v3
Requirements: pip install pandas openpyxl scikit-learn
"""

import pickle
import warnings
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sklearn.preprocessing import StandardScaler, LabelEncoder

warnings.filterwarnings("ignore")

INPUT_CSV   = Path("clinical_preprocessed.csv")
OUTPUT_CSV  = Path("clinical_final.csv")
OUTPUT_XLSX = Path("clinical_final.xlsx")
SCALER_PKL  = Path("scaler.pkl")
ENCODER_PKL = Path("label_encoder.pkl")
CLIP_LOW    = 0.01
CLIP_HIGH   = 0.99
SKIP_OHE    = {"patient_id", "pam50_subtype", "pam50_encoded"}
SKIP_SCALE  = {"pam50_encoded", "survival_event"}


def is_binary(series):
    """Returns True if a Series contains only 0/1 values."""
    vals = set(series.dropna().tolist())
    return vals.issubset({0, 1, 0.0, 1.0})


def load_clinical(path):
    df = pd.read_csv(path)
    print(f"    Loaded {df.shape[0]} patients, {df.shape[1]} columns")
    return df


def engineer_survival(df):
    if "days_to_death" in df.columns and "days_to_last_followup" in df.columns:
        df["overall_survival_days"] = df[["days_to_death", "days_to_last_followup"]].max(axis=1)
    elif "days_to_last_followup" in df.columns:
        df["overall_survival_days"] = df["days_to_last_followup"]
    else:
        df["overall_survival_days"] = None

    if "vital_status" in df.columns:
        df["survival_event"] = df["vital_status"].apply(
            lambda x: 1 if str(x).strip().lower() == "dead" else 0)
    else:
        df["survival_event"] = 0

    print(f"    overall_survival_days missing: {df['overall_survival_days'].isna().sum()}")
    print(f"    survival_event — dead: {(df['survival_event']==1).sum()}, alive: {(df['survival_event']==0).sum()}")
    return df


def onehot_encode_all(df):
    for col in df.select_dtypes(include="bool").columns:
        df[col] = df[col].astype(int)

    obj_cols = [c for c in df.select_dtypes(include="object").columns if c not in SKIP_OHE]

    if not obj_cols:
        print("    No string columns to encode.")
        return df

    print(f"    One-hot encoding {len(obj_cols)} column(s):")
    for c in obj_cols:
        print(f"      {c}: {sorted(df[c].dropna().unique().tolist())}")

    df = pd.get_dummies(df, columns=obj_cols, dummy_na=False)

    for col in df.select_dtypes(include="bool").columns:
        df[col] = df[col].astype(int)

    return df


def clip_outliers(df):
    # Only clip truly continuous columns — explicitly named
    continuous_cols = [c for c in ["age_at_diagnosis", "days_to_last_followup",
                       "overall_survival_days", "tumor_stage_numeric"]
                       if c in df.columns]
    for col in continuous_cols:
        lo = float(df[col].quantile(CLIP_LOW))
        hi = float(df[col].quantile(CLIP_HIGH))
        df[col] = df[col].clip(lower=lo, upper=hi)
    print(f"    Clipped {len(continuous_cols)} continuous columns: {continuous_cols}")
    return df


def scale_features(df):
    scale_cols = [c for c in ["age_at_diagnosis", "days_to_last_followup",
                  "overall_survival_days", "tumor_stage_numeric"]
                  if c in df.columns]
    if not scale_cols:
        print("    No continuous columns to scale.")
        return df, None, []

    print(f"    Scaling {len(scale_cols)} column(s): {scale_cols}")
    scaler = StandardScaler()
    df[scale_cols] = scaler.fit_transform(df[scale_cols])
    return df, scaler, scale_cols


def write_xlsx(df, path):
    summary_rows = [
        ("Total patients",   len(df)),
        ("Total features",   len(df.columns)),
        ("Avg missing rate", f"{df.isna().mean().mean()*100:.1f}%"),
        ("Numeric features", len(df.select_dtypes(include="number").columns)),
        ("PAM50 status",     "Pending — scale up to 50+ patients"),
    ]
    summary_df = pd.DataFrame(summary_rows, columns=["Metric", "Value"])

    with pd.ExcelWriter(str(path), engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Clinical Final", index=False)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)

    wb = load_workbook(str(path))
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    header_fill = PatternFill("solid", start_color="1F4E79")
    alt_fill    = PatternFill("solid", start_color="D6E4F0")
    center      = Alignment(horizontal="center", vertical="center")
    thin        = Side(style="thin", color="CCCCCC")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            fill = alt_fill if row_idx % 2 == 0 else PatternFill()
            for cell in row:
                cell.fill = fill
                cell.alignment = center
                cell.border = border
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value else 0) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 45)
        ws.freeze_panes = "A2"
    wb.save(str(path))


def main():
    print("=" * 65)
    print("  TCGA-BRCA Clinical Preprocessing - Final v3")
    print("=" * 65)

    print("\n[1/6] Loading data ...")
    if not INPUT_CSV.exists():
        print(f"  ERROR: {INPUT_CSV} not found.")
        return
    df = load_clinical(INPUT_CSV)

    print("\n[2/6] Adding PAM50 placeholder columns ...")
    df["pam50_subtype"] = None
    df["pam50_encoded"] = None
    print("    Done (fill when scaling to 50+ patients)")

    print("\n[3/6] Engineering survival columns ...")
    df = engineer_survival(df)

    print("\n[4/6] One-hot encoding categorical columns ...")
    df = onehot_encode_all(df)

    print("\n[5/6] Clipping outliers ...")
    df = clip_outliers(df)

    print("\n[6/6] Scaling continuous features ...")
    df, scaler, scaled_cols = scale_features(df)

    if scaler is not None:
        with open(SCALER_PKL, "wb") as f:
            pickle.dump({"scaler": scaler, "columns": scaled_cols}, f)
        print(f"    Scaler saved -> {SCALER_PKL}")

    le = LabelEncoder()
    with open(ENCODER_PKL, "wb") as f:
        pickle.dump(le, f)
    print(f"    Label encoder saved -> {ENCODER_PKL}")

    print("\n  Exporting ...")
    df.to_csv(OUTPUT_CSV, index=False)
    print(f"    CSV  -> {OUTPUT_CSV.resolve()}")
    write_xlsx(df, OUTPUT_XLSX)
    print(f"    XLSX -> {OUTPUT_XLSX.resolve()}")

    print("\n" + "=" * 65)
    print("  Final columns:")
    print("=" * 65)
    # Deduplicate columns before printing summary
    df = df.loc[:, ~df.columns.duplicated()]
    for col in df.columns:
        flag = " <-- TARGET" if col in ("pam50_subtype", "pam50_encoded") else ""
        dtype = str(df[col].dtype)
        missing = int(df[col].isna().sum())
        print(f"  {col:<52} {dtype:<10} missing={missing}{flag}")

    print(f"\n  Done! {len(df)} patients x {len(df.columns)} features")


if __name__ == "__main__":
    main()
