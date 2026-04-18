#!/usr/bin/env python3
"""
TCGA-BRCA CNV Preprocessing
=============================
Reads Masked Copy Number Segment TSV files (.nocnv_grch38.seg.v2.txt)
and produces a summary feature row per patient (one row per patient).

Each patient folder has 2 files:
  - Smaller file  = Normal sample  (used as baseline reference)
  - Larger file   = Tumor sample   (main signal of interest)

Features extracted per patient (from tumor file):
  Segment-level summary statistics:
    - n_segments              : total number of CNV segments
    - n_gain_segments         : segments with Segment_Mean > 0.2  (amplification)
    - n_loss_segments         : segments with Segment_Mean < -0.2 (deletion)
    - n_neutral_segments      : segments near 0
    - pct_genome_gained       : % of genome with copy gain
    - pct_genome_lost         : % of genome with copy loss
    - mean_segment_mean       : average Segment_Mean across all segments
    - std_segment_mean        : std deviation of Segment_Mean
    - max_segment_mean        : largest amplification
    - min_segment_mean        : largest deletion
    - total_genome_length     : sum of all segment lengths
    - mean_segment_length     : average segment length
    - n_chromosomes_affected  : number of chromosomes with gain or loss

  Per-chromosome Segment_Mean (chr1..chr22, chrX, chrY):
    - chr1_mean, chr2_mean, ... chr22_mean, chrX_mean, chrY_mean

Outputs:
  - cnv_preprocessed.csv
  - cnv_preprocessed.xlsx

Requirements:
  pip install pandas openpyxl scikit-learn
"""

import warnings
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sklearn.preprocessing import StandardScaler

warnings.filterwarnings("ignore")

# ── Configuration ──────────────────────────────────────────────────────────────
DATA_DIR    = Path("TCGA_BRCA_data")
OUTPUT_CSV  = Path("cnv_preprocessed.csv")
OUTPUT_XLSX = Path("cnv_preprocessed.xlsx")

GAIN_THRESHOLD    =  0.2   # Segment_Mean above this = copy gain
LOSS_THRESHOLD    = -0.2   # Segment_Mean below this = copy loss
MISSING_THRESHOLD =  0.60  # Drop column if > 60% missing

CHROMOSOMES = [str(i) for i in range(1, 23)] + ["X", "Y"]
# ──────────────────────────────────────────────────────────────────────────────


def find_cnv_files(data_dir: Path):
    """
    For each patient, find tumor and normal CNV files.
    The larger file is assumed to be the tumor (more segments = more CNV events).
    The smaller file is the normal baseline.
    Returns list of dicts: {patient_id, tumor_file, normal_file}
    """
    patients = []
    for patient_dir in sorted(data_dir.iterdir()):
        cnv_dir = patient_dir / "Copy_Number_Variation"
        if not cnv_dir.exists():
            continue
        files = sorted(cnv_dir.glob("*.txt"), key=lambda f: f.stat().st_size)
        if len(files) == 0:
            print(f"    ⚠ No CNV files for {patient_dir.name}")
            continue
        elif len(files) == 1:
            tumor_file  = files[0]
            normal_file = None
        else:
            normal_file = files[0]   # smaller = normal
            tumor_file  = files[-1]  # larger  = tumor

        patients.append({
            "patient_id"  : patient_dir.name,
            "tumor_file"  : tumor_file,
            "normal_file" : normal_file,
        })
        print(f"    {patient_dir.name}:")
        print(f"      Tumor  → {tumor_file.name}")
        if normal_file:
            print(f"      Normal → {normal_file.name}")
    return patients


def load_seg_file(path: Path) -> pd.DataFrame:
    """Load a .seg.v2.txt file into a DataFrame."""
    df = pd.read_csv(path, sep="\t")
    df.columns = [c.strip() for c in df.columns]
    # Ensure correct types
    df["Chromosome"]    = df["Chromosome"].astype(str).str.replace("chr", "", regex=False)
    df["Start"]         = pd.to_numeric(df["Start"],        errors="coerce")
    df["End"]           = pd.to_numeric(df["End"],          errors="coerce")
    df["Num_Probes"]    = pd.to_numeric(df["Num_Probes"],   errors="coerce")
    df["Segment_Mean"]  = pd.to_numeric(df["Segment_Mean"], errors="coerce")
    df["segment_length"] = df["End"] - df["Start"]
    return df


def extract_features(patient_id: str, tumor_df: pd.DataFrame,
                     normal_df: pd.DataFrame = None) -> dict:
    """
    Extracts a flat feature dict from a tumor segment DataFrame.
    Optionally subtracts normal baseline per chromosome.
    """
    rec = {"patient_id": patient_id}

    seg = tumor_df.copy()

    # If normal file exists, subtract per-chromosome mean as baseline correction
    if normal_df is not None and not normal_df.empty:
        normal_chr_mean = (
            normal_df.groupby("Chromosome")["Segment_Mean"]
            .mean()
            .rename("normal_mean")
        )
        seg = seg.merge(normal_chr_mean, on="Chromosome", how="left")
        seg["Segment_Mean"] = seg["Segment_Mean"] - seg["normal_mean"].fillna(0)
        seg = seg.drop(columns=["normal_mean"])

    # ── Global summary stats ──────────────────────────────────────────────────
    total_len  = seg["segment_length"].sum()
    gain_mask  = seg["Segment_Mean"] >  GAIN_THRESHOLD
    loss_mask  = seg["Segment_Mean"] <  LOSS_THRESHOLD

    rec["n_segments"]             = len(seg)
    rec["n_gain_segments"]        = int(gain_mask.sum())
    rec["n_loss_segments"]        = int(loss_mask.sum())
    rec["n_neutral_segments"]     = int((~gain_mask & ~loss_mask).sum())
    rec["pct_genome_gained"]      = round(
        seg.loc[gain_mask, "segment_length"].sum() / total_len * 100, 4) if total_len > 0 else 0
    rec["pct_genome_lost"]        = round(
        seg.loc[loss_mask, "segment_length"].sum() / total_len * 100, 4) if total_len > 0 else 0
    rec["mean_segment_mean"]      = round(float(seg["Segment_Mean"].mean()), 6)
    rec["std_segment_mean"]       = round(float(seg["Segment_Mean"].std()), 6)
    rec["max_segment_mean"]       = round(float(seg["Segment_Mean"].max()), 6)
    rec["min_segment_mean"]       = round(float(seg["Segment_Mean"].min()), 6)
    rec["total_genome_length"]    = int(total_len)
    rec["mean_segment_length"]    = round(float(seg["segment_length"].mean()), 2)
    rec["n_chromosomes_affected"] = int(
        seg.loc[gain_mask | loss_mask, "Chromosome"].nunique()
    )

    # ── Per-chromosome mean Segment_Mean ─────────────────────────────────────
    chr_means = seg.groupby("Chromosome")["Segment_Mean"].mean()
    for chrom in CHROMOSOMES:
        col = f"chr{chrom}_mean"
        rec[col] = round(float(chr_means[chrom]), 6) if chrom in chr_means.index else None

    return rec


def build_dataframe(records: list) -> pd.DataFrame:
    df = pd.DataFrame(records)
    return df


def drop_high_missing(df: pd.DataFrame) -> pd.DataFrame:
    missing_rate = df.isna().mean()
    to_drop = missing_rate[missing_rate > MISSING_THRESHOLD].index.tolist()
    if to_drop:
        print(f"\n    Dropping {len(to_drop)} column(s) with >{MISSING_THRESHOLD*100:.0f}% missing:")
        for c in to_drop:
            print(f"      • {c}  ({missing_rate[c]*100:.0f}% missing)")
    return df.drop(columns=to_drop)


def impute_missing(df: pd.DataFrame) -> pd.DataFrame:
    """Fill remaining NaN in numeric columns with column median."""
    for col in df.select_dtypes(include="number").columns:
        if df[col].isna().any():
            df[col] = df[col].fillna(df[col].median())
    return df


def scale_features(df: pd.DataFrame) -> tuple:
    """StandardScale all numeric columns except binary-like ones."""
    skip = {"patient_id"}

    def is_binary(series):
        vals = set(series.dropna().tolist())
        return vals.issubset({0, 1, 0.0, 1.0})

    scale_cols = [
        c for c in df.select_dtypes(include="number").columns
        if c not in skip and not is_binary(df[c])
    ]

    if not scale_cols:
        return df, None, []

    scaler = StandardScaler()
    df[scale_cols] = scaler.fit_transform(df[scale_cols])
    print(f"    Scaled {len(scale_cols)} numeric columns.")
    return df, scaler, scale_cols


def write_xlsx(df: pd.DataFrame, path: Path):
    summary_rows = [
        ("Total patients",       len(df)),
        ("Total features",       len(df.columns)),
        ("Avg missing rate",     f"{df.isna().mean().mean()*100:.1f}%"),
        ("CNV gain threshold",   f"Segment_Mean > {GAIN_THRESHOLD}"),
        ("CNV loss threshold",   f"Segment_Mean < {LOSS_THRESHOLD}"),
        ("Chromosomes covered",  len(CHROMOSOMES)),
        ("Baseline correction",  "Normal sample subtracted per chromosome"),
    ]
    summary_df = pd.DataFrame(summary_rows, columns=["Metric", "Value"])

    with pd.ExcelWriter(str(path), engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="CNV Features", index=False)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)

    wb = load_workbook(str(path))
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    header_fill = PatternFill("solid", start_color="1F4E79")
    alt_fill    = PatternFill("solid", start_color="FFF2CC")  # yellow tint for CNV
    center      = Alignment(horizontal="center", vertical="center")
    thin        = Side(style="thin", color="CCCCCC")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            fill = alt_fill if row_idx % 2 == 0 else PatternFill()
            for cell in row:
                cell.fill = fill
                cell.alignment = center
                cell.border = border
        for col in ws.columns:
            max_len = max(
                (len(str(cell.value)) if cell.value is not None else 0)
                for cell in col
            )
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)
        ws.freeze_panes = "A2"
    wb.save(str(path))


def main():
    print("=" * 65)
    print("  TCGA-BRCA CNV Preprocessing")
    print("=" * 65)

    # 1. Find files
    print("\n[1/6] Finding CNV files ...")
    patients = find_cnv_files(DATA_DIR)
    if not patients:
        print("  ERROR: No CNV files found. Run tcga_brca_downloader.py first.")
        return
    print(f"\n    Found {len(patients)} patient(s) with CNV data.")

    # 2. Parse + extract features
    print("\n[2/6] Parsing segments and extracting features ...")
    records = []
    for p in patients:
        try:
            tumor_df  = load_seg_file(p["tumor_file"])
            normal_df = load_seg_file(p["normal_file"]) if p["normal_file"] else None
            rec = extract_features(p["patient_id"], tumor_df, normal_df)
            records.append(rec)
            print(f"    ✓ {p['patient_id']} — {rec['n_segments']} segments, "
                  f"{rec['n_gain_segments']} gains, {rec['n_loss_segments']} losses")
        except Exception as e:
            print(f"    ✗ {p['patient_id']}: {e}")

    if not records:
        print("  No records extracted.")
        return

    # 3. Build DataFrame
    print("\n[3/6] Building DataFrame ...")
    df = build_dataframe(records)
    print(f"    Shape: {df.shape}")

    # 4. Drop high-missing columns
    print("\n[4/6] Dropping high-missing columns ...")
    df = drop_high_missing(df)

    # 5. Impute remaining NaNs
    print("\n[5/6] Imputing remaining NaNs with median ...")
    df = impute_missing(df)
    print(f"    Shape after cleaning: {df.shape}")

    # 6. Scale
    print("\n[6/6] Scaling numeric features (StandardScaler) ...")
    df, scaler, scaled_cols = scale_features(df)

    # Export
    print("\n  Exporting outputs ...")
    df.to_csv(OUTPUT_CSV, index=False)
    print(f"    CSV  saved → {OUTPUT_CSV.resolve()}")
    write_xlsx(df, OUTPUT_XLSX)
    print(f"    XLSX saved → {OUTPUT_XLSX.resolve()}")

    # Summary
    print("\n" + "=" * 65)
    print("  Final column overview:")
    print("=" * 65)
    feature_groups = {
        "Global stats"        : [c for c in df.columns if not c.startswith("chr") and c != "patient_id"],
        "Per-chromosome means": [c for c in df.columns if c.startswith("chr")],
    }
    for group, cols in feature_groups.items():
        print(f"\n  [{group}] ({len(cols)} columns)")
        for col in cols:
            print(f"    {col:<40} {str(df[col].dtype):<10} missing={df[col].isna().sum()}")

    print(f"\n  Done! {len(df)} patient(s) x {len(df.columns)} features")


if __name__ == "__main__":
    main()
